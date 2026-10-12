"""UC-A-30: export data as one JSON document or one Excel workbook."""
from __future__ import annotations

import io
import json
from contextlib import contextmanager
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Participants' voter tokens are set explicitly so the tests can prove they
# never reach an export.
ALICE = "voter-alice-5f1d2c3b4a596877"  # pragma: allowlist secret
BOB = "voter-bob-0a1b2c3d4e5f6071"  # pragma: allowlist secret

SESSION_HEADERS = [
    "Session code", "Title", "Status", "Created at", "Started at", "Ended at",
    "Rating open", "Participant base URL", "Participants", "Questions",
    "Quiz runs", "Feedback responses", "Average rating",
]
PARTICIPANT_HEADERS = [
    "Name", "Email", "Contact number", "Organization", "Role",
    "LinkedIn", "Photo", "Current work", "Learning goal", "First seen", "Last seen",
]
QUESTION_HEADERS = [
    "Question ID", "Asked at", "Asked by", "Question", "Votes",
    "Status", "Responded", "Response", "Responded at", "Removed at",
]
RUN_HEADERS = [
    "Run ID", "Quiz", "Version", "Status", "Questions", "Respondents",
    "Created at", "Started at", "Closed at",
]
ANSWER_HEADERS = [
    "Run ID", "Quiz", "Question no.", "Question", "Respondent",
    "Chosen answers", "Correct", "Answered at",
]
FEEDBACK_HEADERS = [
    "Rating", "Went well", "Could be better",
    "Submitted at", "Updated at",
]
SESSION_SHEETS = [
    "About", "Session", "Participants", "Questions", "Quiz runs", "Quiz answers", "Feedback",
]


def _today() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d")


@contextmanager
def _as_participant(client, token: str):
    admin = client.cookies.get("prashna_admin")
    client.cookies.clear()
    client.cookies.set("voter_token", token)
    try:
        yield
    finally:
        client.cookies.clear()
        if admin:
            client.cookies.set("prashna_admin", admin)


def _sheets(content: bytes) -> dict[str, list[tuple]]:
    wb = load_workbook(io.BytesIO(content))
    return {ws.title: [tuple(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def _sheet_order(content: bytes) -> list[str]:
    return load_workbook(io.BytesIO(content)).sheetnames


def _about(rows: list[tuple]) -> dict:
    return {r[0]: r[1] for r in rows[1:]}


def _records(rows: list[tuple]) -> list[dict]:
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:]]


def _assert_download(r, filename: str, content_type: str) -> None:
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].split(";")[0] == content_type
    assert r.headers["content-disposition"] == f'attachment; filename="{filename}"'
    assert r.headers["cache-control"] == "no-store"


@pytest.fixture
def populated(admin_client, make_session, make_quiz):
    """A live session with two participants, a profile, Q&A, a quiz run
    with answers, and feedback."""
    c = admin_client
    s = make_session("Export Day")
    code = s["code"]
    assert c.post(f"/api/sessions/{code}/start").status_code == 200
    assert c.post(f"/api/sessions/{code}/feedback_enabled", json={"enabled": True}).status_code == 200
    q = make_quiz("Arithmetic")
    run = c.post(f"/api/sessions/{code}/quiz_runs", json={"quiz_id": q["id"]}).json()
    assert c.post(f"/api/sessions/{code}/quiz_runs/{run['id']}/start").status_code == 200

    with _as_participant(c, ALICE):
        # Read the live question as a participant: an admin request to a
        # participant route would add the admin as a participant.
        vq = c.get(f"/api/sessions/{code}/quiz_runs/{run['id']}").json()["current_question"]
        wrong, right = (a["id"] for a in vq["answers"])
        r = c.patch(f"/api/sessions/{code}/me", json={
            "name": "Alice", "email": "alice@example.com", "contact_number": "+1 415 555 0100",
            "organization": "ExampleCo", "role": "Engineer",
            "linkedin_url": "https://www.linkedin.com/in/alice-example",
            "current_work": "Pipelines", "learning_goal": "Testing",
        })
        assert r.status_code == 200, r.text
        alice_id = r.json()["id"]
        query = c.post(f"/api/sessions/{code}/queries", json={"text": "How do exports work?"}).json()
        formula = c.post(f"/api/sessions/{code}/queries", json={"text": "=1+1", "author_name": "Bob"}).json()
        assert c.post(f"/api/sessions/{code}/quiz_runs/{run['id']}/submit",
                      json={"question_id": vq["id"], "answer_ids": [right]}).status_code == 200
        assert c.post(f"/api/sessions/{code}/feedback",
                      json={"rating": 5, "went_well": "Great pace"}).status_code == 200

    with _as_participant(c, BOB):
        bob_id = c.get(f"/api/sessions/{code}/me").json()["id"]
        assert c.post(f"/api/sessions/{code}/queries/{query['id']}/vote").status_code == 200
        assert c.post(f"/api/sessions/{code}/quiz_runs/{run['id']}/submit",
                      json={"question_id": vq["id"], "answer_ids": [wrong]}).status_code == 200
        assert c.post(f"/api/sessions/{code}/feedback",
                      json={"rating": 3, "to_improve": "More breaks"}).status_code == 200

    assert c.post(f"/api/sessions/{code}/queries/{query['id']}/respond",
                  json={"response_text": "As JSON or Excel"}).status_code == 200
    assert c.post(f"/api/sessions/{code}/quiz_runs/{run['id']}/close").status_code == 200
    return {
        "code": code, "mod_key": s["mod_key"], "quiz_id": q["id"], "run_id": run["id"],
        "query_id": query["id"], "formula_id": formula["id"],
        "alice_id": alice_id, "bob_id": bob_id,
    }


def _password_hash() -> str:
    from prashna import models as m
    from prashna.db import SessionLocal
    with SessionLocal() as db:
        return db.query(m.User).filter_by(username="admin").one().password_hash


def _give_to_other_admin(*, session_code: str | None = None, quiz_id: int | None = None) -> None:
    from prashna import auth
    from prashna import models as m
    from prashna.db import SessionLocal
    with SessionLocal() as db:
        other = m.User(username="other_export", password_hash=auth.hash_password("longenough"))
        db.add(other); db.commit(); db.refresh(other)
        if session_code:
            db.query(m.Session).filter_by(code=session_code).one().owner_user_id = other.id
        if quiz_id:
            quiz = db.get(m.Quiz, quiz_id)
            assert quiz is not None
            quiz.owner_user_id = other.id
        db.commit()


# ---------- one session ----------

def test_uc_a_30_session_json(admin_client, populated):
    """UC-A-30: the whole-session JSON nests every table under `session`."""
    code = populated["code"]
    r = admin_client.get(f"/api/sessions/{code}/export?format=json")
    _assert_download(r, f"prashna-session-{code}-{_today()}.json", "application/json")
    body = r.json()
    meta = body["export"]
    assert meta["scope"] == "session"
    assert meta["part"] == "all"
    assert meta["session_code"] == code
    assert meta["exported_at"].endswith("Z")
    assert meta["prashna_version"]
    assert meta["contains_personal_data"] is True
    s = body["session"]
    assert s["code"] == code and s["title"] == "Export Day" and s["status"] == "live"
    assert s["created_at"].endswith("Z")
    assert {"participants", "questions", "quiz_runs", "feedback"} <= set(s)
    assert len(s["participants"]) == 2
    assert all("id" not in p for p in s["participants"])
    alice = next(p for p in s["participants"] if p["name"] == "Alice")
    assert alice["email"] == "alice@example.com"
    assert alice["linkedin_url"] == "https://www.linkedin.com/in/alice-example"
    question = next(q for q in s["questions"] if q["id"] == populated["query_id"])
    assert question["vote_count"] == 1
    assert question["status"] == "responded"
    assert question["response_text"] == "As JSON or Excel"
    assert "participant_id" not in question
    assert question["author_name"] is None
    formula = next(q for q in s["questions"] if q["id"] == populated["formula_id"])
    assert formula["author_name"] == "Bob"
    (run,) = s["quiz_runs"]
    assert run["quiz_title"] == "Arithmetic" and run["version_no"] == 1
    assert run["status"] == "closed" and run["respondent_count"] == 2
    (rq,) = run["questions"]
    assert rq["number"] == 1 and rq["respondents"] == 2
    assert [(a["text"], a["is_correct"], a["count"]) for a in rq["answers"]] == [
        ("3", False, 1), ("4", True, 1)]
    # Respondents are numbered within the run by first answer: Alice, then Bob.
    responses = {x["respondent"]: x for x in run["responses"]}
    assert all("participant_id" not in x and "participant_name" not in x for x in run["responses"])
    assert responses[1]["answers"] == ["4"]
    assert responses[1]["correct"] is True
    assert responses[2]["correct"] is False
    assert responses[1]["question_number"] == 1
    assert all("participant_id" not in f and "author_name" not in f for f in s["feedback"])
    ratings = {f["rating"]: f for f in s["feedback"]}
    assert ratings[5]["went_well"] == "Great pace"
    assert ratings[3]["to_improve"] == "More breaks"


def test_uc_a_30_session_json_is_the_default_format(admin_client, populated):
    """UC-A-30: `format` defaults to json."""
    r = admin_client.get(f"/api/sessions/{populated['code']}/export")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/json")


def test_uc_a_30_session_xlsx(admin_client, populated):
    """UC-A-30: the whole-session workbook has one sheet per table."""
    code = populated["code"]
    r = admin_client.get(f"/api/sessions/{code}/export?format=xlsx")
    _assert_download(r, f"prashna-session-{code}-{_today()}.xlsx", XLSX_TYPE)
    assert _sheet_order(r.content) == SESSION_SHEETS
    sheets = _sheets(r.content)
    assert list(sheets["Session"][0]) == SESSION_HEADERS
    assert list(sheets["Participants"][0]) == PARTICIPANT_HEADERS
    assert list(sheets["Questions"][0]) == QUESTION_HEADERS
    assert list(sheets["Quiz runs"][0]) == RUN_HEADERS
    assert list(sheets["Quiz answers"][0]) == ANSWER_HEADERS
    assert list(sheets["Feedback"][0]) == FEEDBACK_HEADERS

    about = _about(sheets["About"])
    assert about["Export"] == "One session"
    assert about["Session code"] == code
    assert about["Exported at (UTC)"].endswith("Z")
    assert about["Prashna version"]
    assert "personal data" in about["Personal data"]

    (session,) = _records(sheets["Session"])
    assert session["Session code"] == code
    assert session["Participants"] == 2 and session["Questions"] == 2
    assert session["Quiz runs"] == 1 and session["Feedback responses"] == 2
    assert session["Average rating"] == 4
    assert session["Rating open"] == "Yes"
    question = next(q for q in _records(sheets["Questions"]) if q["Question ID"] == populated["query_id"])
    assert question["Votes"] == 1 and question["Responded"] == "Yes"
    assert question["Status"] == "Responded" and question["Asked by"] is None
    answers = {a["Respondent"]: a for a in _records(sheets["Quiz answers"])}
    assert answers[1]["Chosen answers"] == "4"
    assert answers[1]["Correct"] == "Yes"
    assert answers[2]["Correct"] == "No"
    people = {p["Name"]: p for p in _records(sheets["Participants"])}
    assert people["Alice"]["Email"] == "alice@example.com"
    assert sorted(f["Rating"] for f in _records(sheets["Feedback"])) == [3, 5]


@pytest.mark.parametrize("part, scope, keys, sheets", [
    ("profiles", "profiles", {"participants"}, ["About", "Session", "Participants"]),
    ("feedback", "feedback", {"feedback"}, ["About", "Session", "Feedback"]),
    ("qna", "qna", {"questions"}, ["About", "Session", "Questions"]),
    ("quiz_runs", "quiz-runs", {"quiz_runs"}, ["About", "Session", "Quiz runs", "Quiz answers"]),
])
def test_uc_a_30_session_parts(admin_client, populated, part, scope, keys, sheets):
    """UC-A-30: each `part` exports only its own table(s)."""
    code = populated["code"]
    lists = {"participants", "questions", "quiz_runs", "feedback"}
    r = admin_client.get(f"/api/sessions/{code}/export?format=json&part={part}")
    _assert_download(r, f"prashna-{scope}-{code}-{_today()}.json", "application/json")
    body = r.json()
    assert body["export"]["part"] == part
    assert lists & set(body["session"]) == keys
    r = admin_client.get(f"/api/sessions/{code}/export?format=xlsx&part={part}")
    _assert_download(r, f"prashna-{scope}-{code}-{_today()}.xlsx", XLSX_TYPE)
    assert _sheet_order(r.content) == sheets


def test_uc_a_30_profiles_part_lists_only_participants_with_profiles(admin_client, populated):
    """UC-A-30: the profiles export leaves out participants with no profile data."""
    code = populated["code"]
    body = admin_client.get(f"/api/sessions/{code}/export?part=profiles").json()
    assert [p["name"] for p in body["session"]["participants"]] == ["Alice"]
    rows = _sheets(admin_client.get(f"/api/sessions/{code}/export?format=xlsx&part=profiles").content)
    assert [p["Name"] for p in _records(rows["Participants"])] == ["Alice"]


# ---------- all sessions ----------

def test_uc_a_30_all_data(admin_client, populated, make_session, make_quiz):
    """UC-A-30: all owned sessions with all their data, plus the owned quiz
    library; other admins' sessions and quizzes are excluded."""
    mine = make_session("Second")["code"]
    theirs = make_session("Not mine")["code"]
    other_quiz = make_quiz("Their quiz")["id"]
    _give_to_other_admin(session_code=theirs, quiz_id=other_quiz)

    r = admin_client.get("/api/export/all?format=json")
    _assert_download(r, f"prashna-all-data-{_today()}.json", "application/json")
    body = r.json()
    assert body["export"]["scope"] == "all"
    assert body["export"]["session_count"] == 2
    assert {s["code"] for s in body["sessions"]} == {populated["code"], mine}
    full = next(s for s in body["sessions"] if s["code"] == populated["code"])
    assert {"participants", "questions", "quiz_runs", "feedback"} <= set(full)
    quiz_ids = {q["id"] for q in body["quizzes"]}
    assert populated["quiz_id"] in quiz_ids and other_quiz not in quiz_ids
    assert body["export"]["quiz_count"] == len(body["quizzes"])
    exported = next(q for q in body["quizzes"] if q["id"] == populated["quiz_id"])
    assert exported["questions"] and exported["questions"][0]["answers"]

    r = admin_client.get("/api/export/all?format=xlsx")
    _assert_download(r, f"prashna-all-data-{_today()}.xlsx", XLSX_TYPE)
    order = _sheet_order(r.content)
    assert order == ["About", "Sessions", "Participants", "Questions", "Quiz runs", "Quiz answers", "Feedback",
                     "Quizzes", "Quiz questions", "Answer options"]
    sheets = _sheets(r.content)
    assert _about(sheets["About"])["Export"] == "All data"
    assert list(sheets["Sessions"][0]) == SESSION_HEADERS
    for name in order[2:7]:
        assert sheets[name][0][0] == "Session code", name
    for name in order[8:]:
        assert sheets[name][0][0] == "Quiz ID", name
    assert list(sheets["Participants"][0][1:]) == PARTICIPANT_HEADERS
    assert list(sheets["Quiz answers"][0][1:]) == ANSWER_HEADERS
    assert {row[0] for row in sheets["Sessions"][1:]} == {populated["code"], mine}
    assert {row[0] for row in sheets["Questions"][1:]} == {populated["code"]}
    assert populated["quiz_id"] in {row[0] for row in sheets["Quizzes"][1:]}
    assert other_quiz not in {row[0] for row in sheets["Quizzes"][1:]}
    assert any(row[0] == populated["quiz_id"] for row in sheets["Answer options"][1:])


# ---------- one quiz run ----------

def test_uc_a_30_quiz_run(admin_client, populated):
    """UC-A-30: a run export holds the frozen version, results and answers."""
    run_id = populated["run_id"]
    # Renaming the template after the run must not change the run's export.
    admin_client.patch(f"/api/quizzes/{populated['quiz_id']}", json={"title": "Renamed"})
    r = admin_client.get(f"/api/quiz_runs/{run_id}/export?format=json")
    _assert_download(r, f"prashna-quiz-run-{run_id}-{_today()}.json", "application/json")
    body = r.json()
    assert body["export"]["scope"] == "quiz_run"
    assert body["export"]["quiz_run_id"] == run_id
    run = body["quiz_run"]
    assert run["quiz_title"] == "Arithmetic" and run["version_no"] == 1
    assert run["session_code"] == populated["code"]
    assert [(a["text"], a["is_correct"], a["count"]) for a in run["questions"][0]["answers"]] == [
        ("3", False, 1), ("4", True, 1)]
    assert {x["respondent"] for x in run["responses"]} == {1, 2}
    assert body["export"]["contains_personal_data"] is False

    r = admin_client.get(f"/api/quiz_runs/{run_id}/export?format=xlsx")
    _assert_download(r, f"prashna-quiz-run-{run_id}-{_today()}.xlsx", XLSX_TYPE)
    assert _sheet_order(r.content) == ["About", "Quiz run", "Results", "Quiz answers"]
    sheets = _sheets(r.content)
    assert _about(sheets["About"])["Quiz run ID"] == run_id
    assert "Personal data" not in _about(sheets["About"])
    (info,) = _records(sheets["Quiz run"])
    assert info["Quiz"] == "Arithmetic" and info["Version"] == 1 and info["Respondents"] == 2
    results = _records(sheets["Results"])
    assert [(x["Answer"], x["Correct"], x["Count"]) for x in results] == [("3", "No", 1), ("4", "Yes", 1)]
    assert list(sheets["Quiz answers"][0]) == ANSWER_HEADERS


def test_uc_a_30_quiz_run_other_admin_404(admin_client, populated):
    """UC-A-30: another admin's run is not found."""
    _give_to_other_admin(session_code=populated["code"])
    assert admin_client.get(f"/api/quiz_runs/{populated['run_id']}/export").status_code == 404
    assert admin_client.get("/api/quiz_runs/999999/export").status_code == 404


# ---------- one quiz template ----------

def test_uc_a_30_quiz_template(admin_client, populated):
    """UC-A-30: a quiz export holds questions, answers with correctness, and versions."""
    quiz_id = populated["quiz_id"]
    r = admin_client.get(f"/api/quizzes/{quiz_id}/export?format=json")
    _assert_download(r, f"prashna-quiz-{quiz_id}-{_today()}.json", "application/json")
    body = r.json()
    assert body["export"]["scope"] == "quiz"
    assert body["export"]["quiz_id"] == quiz_id
    assert body["export"]["contains_personal_data"] is False
    quiz = body["quiz"]
    assert quiz["title"] == "Arithmetic" and quiz["status"] == "ready"
    (question,) = quiz["questions"]
    assert question["text"] == "2+2?" and question["qtype"] == "single"
    assert [(a["text"], a["is_correct"]) for a in question["answers"]] == [("3", False), ("4", True)]
    assert [(v["version_no"], v["run_count"]) for v in quiz["versions"]] == [(1, 1)]

    r = admin_client.get(f"/api/quizzes/{quiz_id}/export?format=xlsx")
    _assert_download(r, f"prashna-quiz-{quiz_id}-{_today()}.xlsx", XLSX_TYPE)
    assert _sheet_order(r.content) == ["About", "Quiz", "Questions", "Answers", "Versions"]
    sheets = _sheets(r.content)
    about = _about(sheets["About"])
    assert about["Export"] == "One quiz"
    assert "Personal data" not in about
    answers = _records(sheets["Answers"])
    assert [(a["Question no."], a["Answer"], a["Correct"]) for a in answers] == [(1, "3", "No"), (1, "4", "Yes")]
    (version,) = _records(sheets["Versions"])
    assert version["Version"] == 1 and version["Runs"] == 1


def test_uc_a_30_quiz_template_other_admin_404(admin_client, make_quiz):
    """UC-A-30: another admin's quiz is not found."""
    q = make_quiz()
    _give_to_other_admin(quiz_id=q["id"])
    assert admin_client.get(f"/api/quizzes/{q['id']}/export").status_code == 404


# ---------- never exported ----------

def _all_exports(admin_client, populated):
    code = populated["code"]
    urls = ["/api/export/all", f"/api/quiz_runs/{populated['run_id']}/export",
            f"/api/quizzes/{populated['quiz_id']}/export"]
    urls += [f"/api/sessions/{code}/export?part={p}" for p in ("all", "profiles", "feedback", "qna", "quiz_runs")]
    for url in urls:
        sep = "&" if "?" in url else "?"
        for fmt in ("json", "xlsx"):
            r = admin_client.get(f"{url}{sep}format={fmt}")
            assert r.status_code == 200, (url, fmt, r.text)
            yield url, fmt, r


def test_uc_a_30_secrets_never_exported(admin_client, populated):
    """UC-A-30: no voter token, password hash, or mod_key in any export."""
    secrets_ = [ALICE, BOB, _password_hash(), populated["mod_key"]]
    for url, fmt, r in _all_exports(admin_client, populated):
        if fmt == "json":
            text = r.text
            assert "voter_token" not in text and "mod_key" not in text, url
        else:
            cells = [str(v) for rows in _sheets(r.content).values() for row in rows for v in row if v is not None]
            text = "\n".join(cells)
        for secret in secrets_:
            assert secret not in text, (url, fmt)


def test_uc_a_30_exports_never_join_profiles(admin_client, populated):
    """UC-A-30 (Anonymity): the profile name appears only in the Participants table."""
    body = admin_client.get(f"/api/sessions/{populated['code']}/export?format=json").json()["session"]
    assert "Alice" in json.dumps(body["participants"])
    for part in ("questions", "quiz_runs", "feedback"):
        assert "Alice" not in json.dumps(body[part]), part
        assert "participant_id" not in json.dumps(body[part]), part


def test_uc_a_30_xlsx_text_is_never_a_formula(admin_client, populated):
    """UC-A-30: a participant's `=…` entry is a text cell, not a formula."""
    r = admin_client.get(f"/api/sessions/{populated['code']}/export?format=xlsx&part=qna")
    ws = load_workbook(io.BytesIO(r.content))["Questions"]
    cells = [c for row in ws.iter_rows() for c in row if c.value == "=1+1"]
    assert cells and all(c.data_type == "s" for c in cells)


def test_uc_a_30_json_is_valid_utf8_document(admin_client, populated):
    """UC-A-30: the JSON export parses as one document."""
    r = admin_client.get(f"/api/sessions/{populated['code']}/export")
    assert isinstance(json.loads(r.content.decode("utf-8")), dict)


# ---------- access ----------

def test_uc_a_30_delegate_mod_key_refused(client, admin_client, populated, mod_headers):
    """UC-A-30: a delegate moderator (mod_key only) cannot export."""
    code = populated["code"]
    admin_client.cookies.clear()
    r = client.get(f"/api/sessions/{code}/export", headers=mod_headers(populated["mod_key"]))
    assert r.status_code == 401
    assert r.json() == {"detail": "Not authenticated"}


@pytest.mark.parametrize("path", [
    "/api/export/all",
    "/api/sessions/{code}/export",
    "/api/quiz_runs/{run_id}/export",
    "/api/quizzes/{quiz_id}/export",
])
def test_uc_a_30_no_auth_401(client, admin_client, populated, path):
    """UC-A-30: every export needs the admin cookie."""
    admin_client.cookies.clear()
    r = client.get(path.format(code=populated["code"], run_id=populated["run_id"], quiz_id=populated["quiz_id"]))
    assert r.status_code == 401


def test_uc_a_30_other_admins_session_403(admin_client, populated):
    """UC-A-30: another admin's session follows the owner-only convention."""
    _give_to_other_admin(session_code=populated["code"])
    r = admin_client.get(f"/api/sessions/{populated['code']}/export")
    assert r.status_code == 403
    assert r.json() == {"detail": "Not your session"}


def test_uc_a_30_unknown_session_404(admin_client):
    """UC-A-30: an unknown session code is not found."""
    assert admin_client.get("/api/sessions/NOPE99/export").status_code == 404


@pytest.mark.parametrize("path", [
    "/api/export/all?format=csv",
    "/api/sessions/{code}/export?format=csv",
    "/api/quiz_runs/{run_id}/export?format=csv",
    "/api/quizzes/{quiz_id}/export?format=csv",
])
def test_uc_a_30_unknown_format_400(admin_client, populated, path):
    """UC-A-30: an unknown format is refused with a plain message."""
    r = admin_client.get(path.format(code=populated["code"], run_id=populated["run_id"], quiz_id=populated["quiz_id"]))
    assert r.status_code == 400
    assert r.json() == {"detail": "format must be json or xlsx"}


def test_uc_a_30_unknown_part_400(admin_client, populated):
    """UC-A-30: an unknown part is refused with a plain message."""
    r = admin_client.get(f"/api/sessions/{populated['code']}/export?part=everything")
    assert r.status_code == 400
    assert r.json() == {"detail": "part must be one of all, profiles, feedback, qna, quiz_runs"}
